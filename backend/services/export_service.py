import os
from io import BytesIO
from datetime import datetime, date
from PIL import Image as PILImage

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

MONTHS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
DAYS_FR = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

import re

def format_french_date(val, include_day_name=True):
    if not val:
        return ""
    val_str = str(val).strip().rstrip('. ')
    
    # 1. Exact ISO YYYY-MM-DD or DD/MM/YYYY
    dt = None
    if len(val_str) == 10 and val_str[4] == '-' and val_str[7] == '-':
        try:
            dt = datetime.strptime(val_str, '%Y-%m-%d')
        except Exception:
            pass
    elif len(val_str) == 10 and val_str[2] == '/' and val_str[5] == '/':
        try:
            dt = datetime.strptime(val_str, '%d/%m/%Y')
        except Exception:
            pass

    if dt:
        day_name = DAYS_FR[dt.weekday()]
        month_name = MONTHS_FR[dt.month - 1]
        return f"{day_name} {dt.day} {month_name} {dt.year}" if include_day_name else f"{dt.day} {month_name} {dt.year}"

    # 2. Embedded date regex replacement (e.g. 20/09/2026 or 2026-09-20)
    def _repl_slash(match):
        d, m, y = int(match.group(1)), int(match.group(2)), int(match.group(3))
        try:
            dt_emb = datetime(y, m, d)
            d_name = DAYS_FR[dt_emb.weekday()]
            m_name = MONTHS_FR[dt_emb.month - 1]
            return f"{d_name} {d} {m_name} {y}" if include_day_name else f"{d} {m_name} {y}"
        except Exception:
            return match.group(0)

    def _repl_dash(match):
        y, m, d = int(match.group(1)), int(match.group(2)), int(match.group(3))
        try:
            dt_emb = datetime(y, m, d)
            d_name = DAYS_FR[dt_emb.weekday()]
            m_name = MONTHS_FR[dt_emb.month - 1]
            return f"{d_name} {d} {m_name} {y}" if include_day_name else f"{d} {m_name} {y}"
        except Exception:
            return match.group(0)

    res_str = re.sub(r'\b(\d{1,2})/(\d{1,2})/(\d{4})\b', _repl_slash, val_str)
    res_str = re.sub(r'\b(\d{4})-(\d{1,2})-(\d{1,2})\b', _repl_dash, res_str)
    return res_str

class ExportService:

    @staticmethod
    def _clean_val(val, is_date=False):
        if not val:
            return ""
        val_str = str(val).strip().rstrip('. ')
        if is_date:
            return format_french_date(val_str, include_day_name=True)
        return val_str

    # ==========================================
    # WORD (.DOCX) GENERATION (EXACT A4 SPECIFICATION)
    # ==========================================
    @classmethod
    def generate_word(cls, prayer_request, logo_path):
        doc = Document()
        
        # Set A4 Page Dimensions (8.27 x 11.69 inches)
        for section in doc.sections:
            section.page_width = Inches(8.27)
            section.page_height = Inches(11.69)
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Header Table: Logo on Left, Title on Right
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        cell_logo = table.cell(0, 0)
        cell_logo.width = Inches(1.4)
        cell_header = table.cell(0, 1)
        cell_header.width = Inches(5.1)

        # Insert Logo preserving aspect ratio
        if logo_path and os.path.exists(logo_path):
            try:
                p_logo = cell_logo.paragraphs[0]
                run_logo = p_logo.add_run()
                run_logo.add_picture(logo_path, width=Inches(1.1))
            except Exception:
                pass

        # Header text
        p_hdr = cell_header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p_hdr.add_run("EGLISE METHODISTE DE COTE D’IVOIRE\n")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.name = 'Calibri'
        
        r2 = p_hdr.add_run("DISTRICT DE YOPOUGON – CIRCUIT NIANGON\n")
        r2.bold = True
        r2.font.size = Pt(9.5)
        
        r3 = p_hdr.add_run("TEMPLE BETHESDA")
        r3.bold = True
        r3.font.size = Pt(12)
        r3.font.color.rgb = RGBColor(16, 124, 65)

        doc.add_paragraph() # Spacer

        # Title Box
        title_map = {
            'DEMANDE_PRIERE': 'DEMANDE DE PRIERE',
            'NECROLOGIE': 'NECROLOGIE',
            'PRESENTATION_ENFANT': 'PRESENTATION D’ENFANT',
            'REUNION_CLASSE': 'REUNION DE CLASSE METHODISTE'
        }
        title_text = title_map.get(prayer_request.form_type, 'FICHE BETHESDA')

        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run(f"  {title_text}  ")
        t_run.bold = True
        t_run.font.size = Pt(14)
        t_run.font.name = 'Calibri'

        doc.add_paragraph()

        details = prayer_request.details or {}
        fait_a = details.get('fait_a_date', prayer_request.created_at.strftime('%Y-%m-%d') if prayer_request.created_at else '')
        conductor_name = details.get('conducteur', prayer_request.conducteur or 'Conducteur')

        # Build Fields based on Form Type
        if prayer_request.form_type == 'NECROLOGIE':
            cls._add_field(doc, "LA FAMILLE :", details.get('famille', ''))
            cls._add_field(doc, "ET LA CLASSE :", details.get('classe', prayer_request.methode_classe or ''))
            
            p_msg = doc.add_paragraph()
            p_msg.add_run("\nOnt la profonde douleur de vous annoncer le décès,\n").bold = True
            
            sexe_defunt = details.get('sexe_defunt', 'Masculin')
            label_defunt = "DU FRÈRE DÉCÉDÉ :" if sexe_defunt == 'Masculin' else "DE LA SŒUR DÉCÉDÉE :"
            
            cls._add_field(doc, label_defunt, details.get('frere_soeur', ''))
            cls._add_field(doc, "DÉCÉDÉ(E) LE :", details.get('decede_le', ''), is_date=True)
            cls._add_field(doc, "LIEU DE DÉCÈS :", details.get('lieu_deces', ''))
            
            # Veillées
            veillees = details.get('veillees', [])
            if veillees:
                for v in veillees:
                    t = v.get('titre', 'Veillée')
                    l = v.get('lieu_date', '')
                    cls._add_field(doc, f"{t.upper()} :", l, is_date=True)
            else:
                cls._add_field(doc, "LIEU DE VEILLÉE :", details.get('lieu_veillee', ''), is_date=True)

            cls._add_field(doc, "LIEU DE LA LEVÉE :", details.get('lieu_levee', ''))
            cls._add_field(doc, "DATE DE L’ENTERREMENT :", details.get('date_enterrement', ''), is_date=True)
            cls._add_field(doc, "LIEU DE L’ENTERREMENT :", details.get('lieu_enterrement', ''))

            photos = details.get('photos', [])
            cls._add_photos_word(doc, photos)
            cls._add_fait_a(doc, fait_a)
            cls._add_signatures_word(doc, conductor_name)

        elif prayer_request.form_type == 'DEMANDE_PRIERE':
            cls._add_field(doc, "DATE DE LA PRIERE :", details.get('date_priere', prayer_request.event_date.strftime('%Y-%m-%d') if prayer_request.event_date else ''), is_date=True)
            
            doc.add_paragraph()
            # ONLY display requested prayer types in green bold color!
            has_any_prayer = False
            if details.get('priere_soutien'):
                has_any_prayer = True
                p_st = doc.add_paragraph()
                r_st = p_st.add_run("PRIERE DE SOUTIEN :  [ X ]")
                r_st.bold = True
                r_st.font.size = Pt(11)
                r_st.font.color.rgb = RGBColor(16, 124, 65)

            if details.get('priere_guerison'):
                has_any_prayer = True
                p_gr = doc.add_paragraph()
                r_gr = p_gr.add_run("PRIERE DE GUERISON :  [ X ]")
                r_gr.bold = True
                r_gr.font.size = Pt(11)
                r_gr.font.color.rgb = RGBColor(16, 124, 65)

            if details.get('priere_action_grace'):
                has_any_prayer = True
                p_ag = doc.add_paragraph()
                r_ag = p_ag.add_run("PRIERE D’ACTION DE GRACE :  [ X ]")
                r_ag.bold = True
                r_ag.font.size = Pt(11)
                r_ag.font.color.rgb = RGBColor(16, 124, 65)

            if not has_any_prayer:
                p_def = doc.add_paragraph()
                r_def = p_def.add_run("PRIERE D’ACTION DE GRACE :  [ X ]")
                r_def.bold = True
                r_def.font.size = Pt(11)
                r_def.font.color.rgb = RGBColor(16, 124, 65)

            doc.add_paragraph()
            cls._add_field(doc, "CLASSE METHODISTE :", details.get('classe', prayer_request.methode_classe or ''))
            cls._add_field(doc, "CONDUCTEUR (TRICE) :", details.get('conducteur', prayer_request.conducteur or ''))
            cls._add_field(doc, "DEMANDEUR :", details.get('demandeur', prayer_request.demandeur_nom or ''))
            cls._add_field(doc, "SUJET :", details.get('sujet', ''))

            photos = details.get('photos', [])
            cls._add_photos_word(doc, photos)
            cls._add_fait_a(doc, fait_a)
            cls._add_signatures_word(doc, conductor_name)

        elif prayer_request.form_type == 'PRESENTATION_ENFANT':
            cls._add_field(doc, "DATE DE PRÉSENTATION :", details.get('date_presentation', prayer_request.event_date.strftime('%Y-%m-%d') if prayer_request.event_date else ''), is_date=True)
            cls._add_field(doc, "NOM DE L’ENFANT :", details.get('nom_enfant', ''))
            cls._add_field(doc, "GENRE DE L’ENFANT :", details.get('sexe_enfant', details.get('genre_enfant', 'Masculin')))
            cls._add_field(doc, "NOM DU PÈRE :", details.get('nom_pere', ''))
            cls._add_field(doc, "NOM DE LA MÈRE :", details.get('nom_mere', ''))
            cls._add_field(doc, "CLASSE METHODISTE :", details.get('classe', prayer_request.methode_classe or ''))
            cls._add_field(doc, "CONDUCTEUR (TRICE) :", details.get('conducteur', prayer_request.conducteur or ''))

            photos = details.get('photos', [])
            cls._add_photos_word(doc, photos)
            cls._add_fait_a(doc, fait_a)
            cls._add_signatures_word(doc, conductor_name)

        elif prayer_request.form_type == 'REUNION_CLASSE':
            cls._add_field(doc, "Classe méthodiste :", details.get('classe', prayer_request.methode_classe or ''))
            cls._add_field(doc, "Conducteur (trice) :", details.get('conducteur', prayer_request.conducteur or ''))
            cls._add_field(doc, "Date de la réunion :", details.get('date_reunion', prayer_request.event_date.strftime('%Y-%m-%d') if prayer_request.event_date else ''), is_date=True)
            cls._add_field(doc, "Lieu :", details.get('lieu', ''))
            cls._add_field(doc, "Heure :", details.get('heure', ''))
            cls._add_field(doc, "Lieu de rassemblement :", details.get('lieu_rassemblement', ''))

            photos = details.get('photos', [])
            cls._add_photos_word(doc, photos)
            cls._add_fait_a(doc, fait_a)
            cls._add_signatures_word(doc, conductor_name)

        target_io = BytesIO()
        doc.save(target_io)
        target_io.seek(0)
        return target_io

    @classmethod
    def _add_photos_word(cls, doc, photos):
        if not photos:
            return
        p_hdr = doc.add_paragraph()
        r = p_hdr.add_run("Photo(s) Jointe(s) :")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(16, 124, 65)

        for photo_url in photos:
            filename = os.path.basename(photo_url)
            from flask import current_app
            storage_path = os.path.join(current_app.config['STORAGE_DIR'], 'uploads', filename)
            if os.path.exists(storage_path):
                try:
                    with PILImage.open(storage_path) as im:
                        orig_w, orig_h = im.size
                    max_w = 2.2
                    max_h = 2.2
                    ratio = min(max_w / orig_w, max_h / orig_h)
                    w_inch = orig_w * ratio
                    
                    p_img = doc.add_paragraph()
                    p_img.add_run().add_picture(storage_path, width=Inches(w_inch))
                except Exception as e:
                    print(f"Word photo error {filename}: {e}")

    @classmethod
    def _add_field(cls, doc, label, val, is_date=False):
        p = doc.add_paragraph()
        r_lbl = p.add_run(f"{label} ")
        r_lbl.bold = True
        r_val = p.add_run(cls._clean_val(val, is_date=is_date))
        r_val.bold = True
        r_val.font.color.rgb = RGBColor(15, 23, 42)

    @classmethod
    def _add_fait_a(cls, doc, date_str):
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        clean_date = format_french_date(date_str, include_day_name=True)
        r = p.add_run(f"Fait à Abidjan le {clean_date}")
        r.bold = True
        r.italic = True

    @classmethod
    def _add_signatures_word(cls, doc, conductor_name):
        for _ in range(4):
            doc.add_paragraph()
            
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Col 0: Conducteur
        cell_cond = table.cell(0, 0)
        p0 = cell_cond.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run("SIGNATURE DU CONDUCTEUR\n\n")
        r0.bold = True
        r0_sub = p0.add_run(conductor_name)
        r0_sub.italic = True
        r0_sub.font.size = Pt(9.5)

        # Col 1: Bureau du Conseil
        cell_conseil = table.cell(0, 1)
        p1 = cell_conseil.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run("AVIS DU BUREAU DU CONSEIL")
        r1.bold = True

    # ==========================================
    # PDF GENERATION (EXACT A4 SPECIFICATION)
    # ==========================================
    @classmethod
    def generate_pdf(cls, prayer_request, logo_path):
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
        story = []

        styles = getSampleStyleSheet()
        normal_style = styles['Normal']
        
        hdr_style = ParagraphStyle(
            'HeaderStyle',
            parent=normal_style,
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#0F172A')
        )
        
        title_box_style = ParagraphStyle(
            'TitleBoxStyle',
            parent=normal_style,
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            alignment=1, # Center
            textColor=colors.HexColor('#107C41')
        )
        
        field_style = ParagraphStyle(
            'FieldStyle',
            parent=normal_style,
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=16,
            textColor=colors.HexColor('#0F172A')
        )

        field_bold = ParagraphStyle(
            'FieldBold',
            parent=normal_style,
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=16,
            textColor=colors.HexColor('#0F172A')
        )

        prayer_highlight_style = ParagraphStyle(
            'PrayerHighlight',
            parent=normal_style,
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=18,
            textColor=colors.HexColor('#107C41')
        )

        # Logo and Header
        logo_img = ""
        if logo_path and os.path.exists(logo_path):
            try:
                logo_img = RLImage(logo_path, width=70, height=70)
            except Exception:
                logo_img = ""
        
        hdr_text = Paragraph(
            "<b>EGLISE METHODISTE DE COTE D’IVOIRE</b><br/>"
            "DISTRICT DE YOPOUGON – CIRCUIT NIANGON<br/>"
            "<font color='#107C41'><b>TEMPLE BETHESDA</b></font>",
            hdr_style
        )
        
        if logo_img:
            tbl_hdr = Table([[logo_img, hdr_text]], colWidths=[90, 430])
        else:
            tbl_hdr = Table([[hdr_text]], colWidths=[520])

        tbl_hdr.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ]))
        story.append(tbl_hdr)
        story.append(Spacer(1, 15))

        # Title Box
        title_map = {
            'DEMANDE_PRIERE': 'DEMANDE DE PRIERE',
            'NECROLOGIE': 'NECROLOGIE',
            'PRESENTATION_ENFANT': 'PRESENTATION D’ENFANT',
            'REUNION_CLASSE': 'REUNION DE CLASSE METHODISTE'
        }
        title_text = title_map.get(prayer_request.form_type, 'FICHE BETHESDA')
        
        tbl_title = Table([[Paragraph(f"<b>{title_text}</b>", title_box_style)]], colWidths=[400])
        tbl_title.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor('#107C41')),
            ('PADDING', (0,0), (-1,-1), 8),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))
        story.append(tbl_title)
        story.append(Spacer(1, 15))

        details = prayer_request.details or {}
        fait_a = details.get('fait_a_date', prayer_request.created_at.strftime('%Y-%m-%d') if prayer_request.created_at else '')
        conductor_name = details.get('conducteur', prayer_request.conducteur or 'Conducteur')

        # Content based on form_type
        if prayer_request.form_type == 'NECROLOGIE':
            cls._pdf_add_field(story, "LA FAMILLE :", details.get('famille', ''), field_bold, field_style)
            cls._pdf_add_field(story, "ET LA CLASSE :", details.get('classe', prayer_request.methode_classe or ''), field_bold, field_style)
            story.append(Spacer(1, 8))
            story.append(Paragraph("<b>Ont la profonde douleur de vous annoncer le décès,</b>", field_bold))
            story.append(Spacer(1, 8))
            
            sexe_defunt = details.get('sexe_defunt', 'Masculin')
            label_defunt = "DU FRÈRE DÉCÉDÉ :" if sexe_defunt == 'Masculin' else "DE LA SŒUR DÉCÉDÉE :"
            
            cls._pdf_add_field(story, label_defunt, details.get('frere_soeur', ''), field_bold, field_style)
            cls._pdf_add_field(story, "DÉCÉDÉ(E) LE :", details.get('decede_le', ''), field_bold, field_style, is_date=True)
            cls._pdf_add_field(story, "LIEU DE DÉCÈS :", details.get('lieu_deces', ''), field_bold, field_style)
            
            # Veillées
            veillees = details.get('veillees', [])
            if veillees:
                for v in veillees:
                    t = v.get('titre', 'Veillée')
                    l = v.get('lieu_date', '')
                    cls._pdf_add_field(story, f"{t.upper()} :", l, field_bold, field_style, is_date=True)
            else:
                cls._pdf_add_field(story, "LIEU DE VEILLÉE :", details.get('lieu_veillee', ''), field_bold, field_style, is_date=True)

            cls._pdf_add_field(story, "LIEU DE LA LEVÉE :", details.get('lieu_levee', ''), field_bold, field_style)
            cls._pdf_add_field(story, "DATE DE L’ENTERREMENT :", details.get('date_enterrement', ''), field_bold, field_style, is_date=True)
            cls._pdf_add_field(story, "LIEU DE L’ENTERREMENT :", details.get('lieu_enterrement', ''), field_bold, field_style)
            
            photos = details.get('photos', [])
            cls._pdf_add_photos(story, photos, field_bold)
            cls._pdf_add_fait_a(story, fait_a, field_bold)
            cls._pdf_add_signatures_pdf(story, conductor_name, field_bold)

        elif prayer_request.form_type == 'DEMANDE_PRIERE':
            cls._pdf_add_field(story, "DATE DE LA PRIERE :", details.get('date_priere', prayer_request.event_date.strftime('%Y-%m-%d') if prayer_request.event_date else ''), field_bold, field_style, is_date=True)
            story.append(Spacer(1, 8))
            
            # ONLY display requested prayer types in green bold font!
            has_any_prayer = False
            if details.get('priere_soutien'):
                has_any_prayer = True
                story.append(Paragraph("<b>PRIERE DE SOUTIEN :  [ X ]</b>", prayer_highlight_style))
                story.append(Spacer(1, 4))
            if details.get('priere_guerison'):
                has_any_prayer = True
                story.append(Paragraph("<b>PRIERE DE GUERISON :  [ X ]</b>", prayer_highlight_style))
                story.append(Spacer(1, 4))
            if details.get('priere_action_grace'):
                has_any_prayer = True
                story.append(Paragraph("<b>PRIERE D’ACTION DE GRACE :  [ X ]</b>", prayer_highlight_style))
                story.append(Spacer(1, 4))

            if not has_any_prayer:
                story.append(Paragraph("<b>PRIERE D’ACTION DE GRACE :  [ X ]</b>", prayer_highlight_style))
                story.append(Spacer(1, 4))

            story.append(Spacer(1, 8))
            cls._pdf_add_field(story, "CLASSE METHODISTE :", details.get('classe', prayer_request.methode_classe or ''), field_bold, field_style)
            cls._pdf_add_field(story, "CONDUCTEUR (TRICE) :", details.get('conducteur', prayer_request.conducteur or ''), field_bold, field_style)
            cls._pdf_add_field(story, "DEMANDEUR :", details.get('demandeur', prayer_request.demandeur_nom or ''), field_bold, field_style)
            cls._pdf_add_field(story, "SUJET :", details.get('sujet', ''), field_bold, field_style)
            
            photos = details.get('photos', [])
            cls._pdf_add_photos(story, photos, field_bold)
            cls._pdf_add_fait_a(story, fait_a, field_bold)
            cls._pdf_add_signatures_pdf(story, conductor_name, field_bold)

        elif prayer_request.form_type == 'PRESENTATION_ENFANT':
            cls._pdf_add_field(story, "DATE DE PRÉSENTATION :", details.get('date_presentation', prayer_request.event_date.strftime('%Y-%m-%d') if prayer_request.event_date else ''), field_bold, field_style, is_date=True)
            cls._pdf_add_field(story, "NOM DE L’ENFANT :", details.get('nom_enfant', ''), field_bold, field_style)
            cls._pdf_add_field(story, "GENRE DE L’ENFANT :", details.get('sexe_enfant', details.get('genre_enfant', 'Masculin')), field_bold, field_style)
            cls._pdf_add_field(story, "NOM DU PÈRE :", details.get('nom_pere', ''), field_bold, field_style)
            cls._pdf_add_field(story, "NOM DE LA MÈRE :", details.get('nom_mere', ''), field_bold, field_style)
            cls._pdf_add_field(story, "CLASSE METHODISTE :", details.get('classe', prayer_request.methode_classe or ''), field_bold, field_style)
            cls._pdf_add_field(story, "CONDUCTEUR (TRICE) :", details.get('conducteur', prayer_request.conducteur or ''), field_bold, field_style)
            
            photos = details.get('photos', [])
            cls._pdf_add_photos(story, photos, field_bold)
            cls._pdf_add_fait_a(story, fait_a, field_bold)
            cls._pdf_add_signatures_pdf(story, conductor_name, field_bold)

        elif prayer_request.form_type == 'REUNION_CLASSE':
            cls._pdf_add_field(story, "Classe méthodiste :", details.get('classe', prayer_request.methode_classe or ''), field_bold, field_style)
            cls._pdf_add_field(story, "Conducteur (trice) :", details.get('conducteur', prayer_request.conducteur or ''), field_bold, field_style)
            cls._pdf_add_field(story, "Date de la réunion :", details.get('date_reunion', prayer_request.event_date.strftime('%Y-%m-%d') if prayer_request.event_date else ''), field_bold, field_style, is_date=True)
            cls._pdf_add_field(story, "Lieu :", details.get('lieu', ''), field_bold, field_style)
            cls._pdf_add_field(story, "Heure :", details.get('heure', ''), field_bold, field_style)
            cls._pdf_add_field(story, "Lieu de rassemblement :", details.get('lieu_rassemblement', ''), field_bold, field_style)
            
            photos = details.get('photos', [])
            cls._pdf_add_photos(story, photos, field_bold)
            cls._pdf_add_fait_a(story, fait_a, field_bold)
            cls._pdf_add_signatures_pdf(story, conductor_name, field_bold)

        doc.build(story)
        buffer.seek(0)
        return buffer

    @classmethod
    def _pdf_add_photos(cls, story, photos, style_bold):
        if not photos:
            return
        story.append(Spacer(1, 10))
        story.append(Paragraph("<b>Photo(s) Jointe(s) :</b>", style_bold))
        story.append(Spacer(1, 5))
        img_cells = []
        for photo_url in photos:
            filename = os.path.basename(photo_url)
            from flask import current_app
            storage_path = os.path.join(current_app.config['STORAGE_DIR'], 'uploads', filename)
            if os.path.exists(storage_path):
                try:
                    with PILImage.open(storage_path) as im:
                        orig_w, orig_h = im.size
                    max_w = 180
                    max_h = 140
                    ratio = min(max_w / orig_w, max_h / orig_h)
                    w = int(orig_w * ratio)
                    h = int(orig_h * ratio)
                    img_cells.append(RLImage(storage_path, width=w, height=h))
                except Exception as e:
                    print(f"Skipping unreadable PDF image {filename}: {e}")
        if img_cells:
            t = Table([img_cells], colWidths=[200] * len(img_cells))
            t.setStyle(TableStyle([
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

    @classmethod
    def _pdf_add_field(cls, story, label, val, style_bold, style_normal, is_date=False):
        clean = cls._clean_val(val, is_date=is_date)
        p = Paragraph(f"<b>{label}</b> &nbsp; {clean}", style_normal)
        story.append(p)
        story.append(Spacer(1, 6))

    @classmethod
    def _pdf_add_fait_a(cls, story, fait_a_str, style_bold):
        story.append(Spacer(1, 20))
        clean_date = format_french_date(fait_a_str, include_day_name=True)
        p = Paragraph(f"<font size=10><b><i>Fait à Abidjan le {clean_date}</i></b></font>", ParagraphStyle('RightText', parent=style_bold, alignment=2))
        story.append(p)
        story.append(Spacer(1, 35))

    @classmethod
    def _pdf_add_signatures_pdf(cls, story, conductor_name, style_bold):
        p_cond = Paragraph(f"<b>SIGNATURE DU CONDUCTEUR</b><br/><br/><i><font size=8>{conductor_name}</font></i>", ParagraphStyle('CenterSig1', parent=style_bold, alignment=1, fontSize=8.5))
        p_bureau = Paragraph("<b>AVIS DU BUREAU DU CONSEIL</b>", ParagraphStyle('CenterSig2', parent=style_bold, alignment=1, fontSize=8.5))
        
        t = Table([[p_cond, p_bureau]], colWidths=[257, 257])
        t.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        story.append(KeepTogether(t))

